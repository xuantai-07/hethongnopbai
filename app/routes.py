import io
import re
import tempfile
import uuid
import zipfile
from datetime import datetime
from pathlib import Path

from flask import (
    Blueprint, render_template, request, session,
    redirect, url_for, flash, current_app, send_from_directory,
    send_file,
)
from flask_login import login_user, logout_user, login_required, current_user
from sqlalchemy import func
from sqlalchemy.orm import joinedload

from app.helpers import clean_display_filename, get_extension, rate_limiter
from app.models import db, Teacher, Assignment, Submission

main = Blueprint("main", __name__)

# ── Định dạng tệp được phép ──────────────────────────────────────

ALLOWED_EXTENSIONS = {
    "pdf", "doc", "docx", "xls", "xlsx",
    "ppt", "pptx", "txt", "zip", "rar",
    "py", "cpp", "c", "java", "html", "css", "js",
}


def allowed_file(filename):
    return get_extension(filename) in ALLOWED_EXTENSIONS


ALLOWED_ASSIGNMENT_EXTENSIONS = {
    "pdf", "doc", "docx", "ppt", "pptx", "txt", "zip",
}


def allowed_assignment_file(filename):
    return get_extension(filename) in ALLOWED_ASSIGNMENT_EXTENSIONS


# Số báo danh / mã sinh viên: chữ không dấu, số, dấu chấm, gạch ngang, gạch dưới
STUDENT_ID_PATTERN = re.compile(r"^[A-Za-z0-9._-]{2,50}$")
STUDENT_NAME_MAX_LENGTH = 150


def client_ip():
    return request.remote_addr or "unknown"


# ── Helpers dùng chung ────────────────────────────────────────────

def _get_open_assignments():
    """Lấy bài tập đang mở VÀ chưa quá hạn (nhóm 1.1)."""
    now = datetime.now()
    return (
        Assignment.query
        .filter_by(is_open=True)
        
        .filter(db.or_(
            Assignment.deadline.is_(None),
            Assignment.deadline >= now,
        ))
        .order_by(Assignment.created_at.desc())
        .all()
    )


def _render_submit_form(form_data=None):
    """Render trang nộp bài, tuỳ chọn giữ lại dữ liệu form (nhóm 1.2)."""
    assignments = _get_open_assignments()
    exts = sorted(ALLOWED_EXTENSIONS)
    max_bytes = current_app.config["MAX_CONTENT_LENGTH"]
    return render_template(
        "submit.html",
        assignments=assignments,
        form_data=form_data or {},
        allowed_extensions_display=", ".join(ext.upper() for ext in exts),
        allowed_extensions_accept=",".join(f".{ext}" for ext in exts),
        allowed_extensions_list=",".join(exts),
        max_size_mb=max_bytes // (1024 * 1024),
        max_content_length=max_bytes,
    )


def _latest_submission_ids(assignment_id=None):
    """Tập ID bản nộp mới nhất của mỗi sinh viên mỗi bài (dùng MAX(id))."""
    q = db.session.query(func.max(Submission.id))
    if assignment_id:
        q = q.filter(Submission.assignment_id == assignment_id)
    rows = q.group_by(Submission.assignment_id, Submission.student_code).all()
    return {row[0] for row in rows}


def _sanitize_excel_value(val):
    """Chống lỗi Formula Injection khi mở tệp Excel."""
    if isinstance(val, str) and val.startswith(("=", "+", "-", "@", "\t", "\r")):
        return f"'{val}"
    return val


# =====================================================================
#  TRANG SINH VIÊN
# =====================================================================

@main.route("/")
def index():
    return _render_submit_form()


@main.route("/submit", methods=["POST"])
def submit_assignment():
    assignment_id = request.form.get("assignment_id", type=int)
    student_name = " ".join(request.form.get("student_name", "").split())
    student_number = request.form.get("student_number", "").strip().upper()
    student_code = request.form.get("student_code", "").strip().upper()
    uploaded_file = request.files.get("file")

    # Nhóm 1.2: giữ lại dữ liệu khi nộp lỗi
    form_data = {
        "assignment_id": assignment_id,
        "student_name": student_name,
        "student_number": request.form.get("student_number", "").strip(),
        "student_code": request.form.get("student_code", "").strip(),
    }

    if not all([assignment_id, student_name, student_number, student_code]):
        flash("Vui lòng nhập đầy đủ thông tin.", "error")
        return _render_submit_form(form_data)

    if len(student_name) > STUDENT_NAME_MAX_LENGTH:
        flash(
            f"Họ và tên tối đa {STUDENT_NAME_MAX_LENGTH} ký tự.", "error"
        )
        return _render_submit_form(form_data)

    if not STUDENT_ID_PATTERN.match(student_number):
        flash(
            "Số báo danh chỉ gồm chữ không dấu, số, dấu chấm, "
            "gạch ngang hoặc gạch dưới (2–50 ký tự).",
            "error",
        )
        return _render_submit_form(form_data)

    if not STUDENT_ID_PATTERN.match(student_code):
        flash(
            "Mã sinh viên chỉ gồm chữ không dấu, số, dấu chấm, "
            "gạch ngang hoặc gạch dưới (2–50 ký tự).",
            "error",
        )
        return _render_submit_form(form_data)

    assignment = db.session.get(Assignment, assignment_id)

    if not assignment or not assignment.is_open:
        flash("Bài tập không tồn tại hoặc đã đóng.", "error")
        return _render_submit_form(form_data)

    if not uploaded_file or not uploaded_file.filename:
        flash("Vui lòng chọn tệp cần nộp.", "error")
        return _render_submit_form(form_data)

    # Giữ tên gốc (có dấu tiếng Việt) để hiển thị; kiểm tra đuôi SAU khi làm sạch
    original_filename = clean_display_filename(uploaded_file.filename)
    extension = get_extension(original_filename)

    if not original_filename:
        flash("Tên tệp không hợp lệ.", "error")
        return _render_submit_form(form_data)

    # Nhóm 1.5: ghi rõ định dạng cho phép khi từ chối
    if not allowed_file(original_filename):
        exts = ", ".join(sorted(ext.upper() for ext in ALLOWED_EXTENSIONS))
        flash(
            f"Định dạng tệp chưa được hỗ trợ. Các định dạng cho phép: {exts}.",
            "error",
        )
        return _render_submit_form(form_data)

    if assignment.deadline and datetime.now() > assignment.deadline:
        flash("Đã quá hạn nộp bài.", "error")
        return _render_submit_form(form_data)

    # Chống spam: giới hạn số lần nộp theo IP và theo sinh viên + bài tập
    config = current_app.config
    window = config["SUBMIT_LIMIT_WINDOW"]
    ip_key = f"submit-ip:{client_ip()}"
    student_key = f"submit-student:{assignment.id}:{student_code}"

    if (
        rate_limiter.is_limited(ip_key, config["SUBMIT_LIMIT_PER_IP"], window)
        or rate_limiter.is_limited(
            student_key, config["SUBMIT_LIMIT_PER_STUDENT"], window
        )
    ):
        flash(
            "Bạn đã nộp quá nhiều lần trong thời gian ngắn. "
            "Vui lòng thử lại sau ít phút.",
            "error",
        )
        return _render_submit_form(form_data)

    # Tên lưu trên đĩa chỉ gồm UUID + đuôi tệp, không phụ thuộc tên gốc
    stored_filename = f"{uuid.uuid4().hex}.{extension}"
    upload_dir = Path(current_app.config["UPLOAD_FOLDER"])
    upload_dir.mkdir(parents=True, exist_ok=True)
    saved_path = upload_dir / stored_filename

    try:
        uploaded_file.save(saved_path)

        if saved_path.stat().st_size == 0:
            saved_path.unlink(missing_ok=True)
            flash("Tệp bài làm đang trống (0 byte). Vui lòng kiểm tra lại.", "error")
            return _render_submit_form(form_data)

        submission = Submission(
            assignment_id=assignment.id,
            student_name=student_name,
            student_number=student_number,
            student_code=student_code,
            original_filename=original_filename,
            stored_filename=stored_filename,
        )

        db.session.add(submission)
        db.session.commit()

    except Exception:
        db.session.rollback()

        if saved_path.exists():
            saved_path.unlink()

        current_app.logger.exception("Không thể lưu bài nộp.")
        flash("Có lỗi khi lưu bài nộp. Vui lòng thử lại.", "error")
        return _render_submit_form(form_data)

    rate_limiter.hit(ip_key, window)
    rate_limiter.hit(student_key, window)

    # Post/Redirect/Get: tải lại trang thành công sẽ KHÔNG nộp lại bài
    session["last_submission_id"] = submission.id
    return redirect(url_for("main.submission_success"))


@main.route("/submit/success")
def submission_success():
    submission_id = session.get("last_submission_id")
    submission = (
        db.session.get(Submission, submission_id) if submission_id else None
    )

    if submission is None:
        return redirect(url_for("main.index"))

    return render_template("success.html", submission=submission)


# =====================================================================
#  ĐĂNG NHẬP / ĐĂNG XUẤT GIẢNG VIÊN
# =====================================================================

@main.route("/teacher/login", methods=["GET", "POST"])
def teacher_login():
    if current_user.is_authenticated:
        return redirect(url_for("main.teacher_dashboard"))

    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")

        # Nhóm 1.4: giới hạn số lần đăng nhập sai
        config = current_app.config
        login_key = f"login-ip:{client_ip()}"
        login_window = config.get("LOGIN_LIMIT_WINDOW", 600)
        login_limit = config.get("LOGIN_LIMIT_PER_IP", 5)

        if rate_limiter.is_limited(login_key, login_limit, login_window):
            flash(
                "Quá nhiều lần đăng nhập sai. "
                "Vui lòng thử lại sau ít phút.",
                "error",
            )
            return render_template("login.html")

        teacher = Teacher.query.filter_by(username=username).first()

        if teacher and teacher.check_password(password):
            rate_limiter.reset(login_key)
            login_user(teacher)
            return redirect(url_for("main.teacher_dashboard"))

        rate_limiter.hit(login_key, login_window)
        flash("Tên đăng nhập hoặc mật khẩu không đúng.", "error")

    return render_template("login.html")


@main.route("/teacher/logout", methods=["POST"])
@login_required
def teacher_logout():
    logout_user()
    flash("Bạn đã đăng xuất.", "success")
    return redirect(url_for("main.teacher_login"))


@main.route("/teacher/change-password", methods=["POST"])
@login_required
def change_password():
    current_password = request.form.get("current_password", "")
    new_password = request.form.get("new_password", "")
    confirm_password = request.form.get("confirm_password", "")

    if not current_user.check_password(current_password):
        flash("Mật khẩu hiện tại không đúng.", "error")
        return redirect(url_for("main.teacher_dashboard"))

    if len(new_password) < 12:
        flash("Mật khẩu mới phải có ít nhất 12 ký tự.", "error")
        return redirect(url_for("main.teacher_dashboard"))

    if new_password != confirm_password:
        flash("Mật khẩu xác nhận không khớp.", "error")
        return redirect(url_for("main.teacher_dashboard"))

    try:
        teacher = db.session.get(Teacher, current_user.id)
        teacher.set_password(new_password)
        db.session.commit()
        flash("Đã đổi mật khẩu thành công.", "success")
    except Exception:
        db.session.rollback()
        current_app.logger.exception("Không thể đổi mật khẩu.")
        flash("Có lỗi khi đổi mật khẩu.", "error")

    return redirect(url_for("main.teacher_dashboard"))


# =====================================================================
#  BẢNG ĐIỀU KHIỂN GIẢNG VIÊN
# =====================================================================

@main.route("/teacher")
@login_required
def teacher_dashboard():
    # Danh sách bài tập
    assignments = Assignment.query.order_by(
        Assignment.created_at.desc()
    ).all()

    # Nhóm 2.2: Tính số lượt nộp qua GROUP BY thay vì N+1
    count_rows = (
        db.session.query(
            Submission.assignment_id,
            func.count(Submission.id),
        )
        .group_by(Submission.assignment_id)
        .all()
    )
    submission_counts = dict(count_rows)

    # Nhóm 2.3: phân trang + lọc theo bài tập
    page = request.args.get("page", 1, type=int)
    filter_assignment_id = request.args.get("assignment_id", 0, type=int)

    # Đánh dấu bản nộp mới nhất (tối ưu phạm vi theo bộ lọc hiện tại)
    latest_ids = _latest_submission_ids(filter_assignment_id or None)

    # Nhóm 2.2: joinedload tránh N+1 khi hiển thị tên bài tập
    query = Submission.query.options(joinedload(Submission.assignment))

    if filter_assignment_id:
        query = query.filter(Submission.assignment_id == filter_assignment_id)

    per_page = current_app.config.get("SUBMISSIONS_PER_PAGE", 20)
    pagination = query.order_by(
        Submission.submitted_at.desc(),
        Submission.id.desc(),
    ).paginate(page=page, per_page=per_page, error_out=False)

    submissions = pagination.items

    return render_template(
        "teacher.html",
        assignments=assignments,
        submissions=submissions,
        latest_submission_ids=latest_ids,
        submission_counts=submission_counts,
        pagination=pagination,
        filter_assignment_id=filter_assignment_id,
    )


# =====================================================================
#  QUẢN LÝ BÀI TẬP
# =====================================================================

@main.route("/teacher/assignments/create", methods=["POST"])
@login_required
def create_assignment():
    title = request.form.get("title", "").strip()
    description = request.form.get("description", "").strip()
    deadline_text = request.form.get("deadline", "").strip()

    assignment_file = request.files.get("assignment_file")

    if not title:
        flash("Vui lòng nhập tên bài tập.", "error")
        return redirect(url_for("main.teacher_dashboard"))

    # Bắt buộc phải tải lên tệp đề bài
    if not assignment_file or not assignment_file.filename:
        flash("Vui lòng chọn tệp đề bài.", "error")
        return redirect(url_for("main.teacher_dashboard"))

    if not allowed_assignment_file(assignment_file.filename):
        flash(
            "Định dạng đề bài không hợp lệ. "
            "Hãy sử dụng PDF, DOC, DOCX, PPT, PPTX, TXT hoặc ZIP.",
            "error",
        )
        return redirect(url_for("main.teacher_dashboard"))

    deadline = None

    if deadline_text:
        try:
            deadline = datetime.strptime(
                deadline_text, "%Y-%m-%dT%H:%M"
            )
        except ValueError:
            flash("Thời hạn nộp không hợp lệ.", "error")
            return redirect(url_for("main.teacher_dashboard"))

    original_filename = clean_display_filename(assignment_file.filename)

    if not original_filename or not allowed_assignment_file(original_filename):
        flash("Tên tệp đề bài không hợp lệ.", "error")
        return redirect(url_for("main.teacher_dashboard"))

    stored_filename = (
        f"{uuid.uuid4().hex}.{get_extension(original_filename)}"
    )

    assignment_dir = Path(current_app.config["ASSIGNMENT_FOLDER"])
    assignment_dir.mkdir(parents=True, exist_ok=True)

    saved_path = assignment_dir / stored_filename

    try:
        assignment_file.save(saved_path)

        assignment = Assignment(
            title=title,
            description=description,
            assignment_original_filename=original_filename,
            assignment_stored_filename=stored_filename,
            deadline=deadline,
            is_open=True,
        )

        db.session.add(assignment)
        db.session.commit()

    except Exception:
        db.session.rollback()

        if saved_path.exists():
            saved_path.unlink()

        current_app.logger.exception("Không thể tạo bài tập.")
        flash("Có lỗi khi tạo bài tập. Vui lòng thử lại.", "error")
        return redirect(url_for("main.teacher_dashboard"))

    flash("Đã tạo bài tập và tải lên đề bài thành công.", "success")
    return redirect(url_for("main.teacher_dashboard"))


@main.route(
    "/teacher/assignments/<int:assignment_id>/toggle", methods=["POST"]
)
@login_required
def toggle_assignment(assignment_id):
    assignment = db.session.get(Assignment, assignment_id)

    if assignment is None:
        flash("Không tìm thấy bài tập.", "error")
        return redirect(url_for("main.teacher_dashboard"))

    assignment.is_open = not assignment.is_open
    db.session.commit()

    status = "mở lại" if assignment.is_open else "đóng"
    flash(f"Đã {status} bài tập «{assignment.title}».", "success")
    return redirect(url_for("main.teacher_dashboard"))


# ── Nhóm 3.3: Sửa bài tập ──

@main.route(
    "/teacher/assignments/<int:assignment_id>/edit", methods=["POST"]
)
@login_required
def edit_assignment(assignment_id):
    assignment = db.session.get(Assignment, assignment_id)
    if assignment is None:
        flash("Không tìm thấy bài tập.", "error")
        return redirect(url_for("main.teacher_dashboard"))

    title = request.form.get("title", "").strip()
    description = request.form.get("description", "").strip()
    deadline_text = request.form.get("deadline", "").strip()
    assignment_file = request.files.get("assignment_file")

    if not title:
        flash("Vui lòng nhập tên bài tập.", "error")
        return redirect(url_for("main.teacher_dashboard"))

    deadline = None
    if deadline_text:
        try:
            deadline = datetime.strptime(deadline_text, "%Y-%m-%dT%H:%M")
        except ValueError:
            flash("Thời hạn nộp không hợp lệ.", "error")
            return redirect(url_for("main.teacher_dashboard"))

    assignment.title = title
    assignment.description = description
    assignment.deadline = deadline

    old_stored_filename = None
    new_saved_path = None
    assignment_dir = Path(current_app.config["ASSIGNMENT_FOLDER"])

    # Thay tệp đề bài nếu giảng viên tải lên tệp mới
    if assignment_file and assignment_file.filename:
        if not allowed_assignment_file(assignment_file.filename):
            flash(
                "Định dạng đề bài không hợp lệ. "
                "Hãy sử dụng PDF, DOC, DOCX, PPT, PPTX, TXT hoặc ZIP.",
                "error",
            )
            return redirect(url_for("main.teacher_dashboard"))

        original_filename = clean_display_filename(assignment_file.filename)
        if not original_filename or not allowed_assignment_file(original_filename):
            flash("Tên tệp đề bài không hợp lệ.", "error")
            return redirect(url_for("main.teacher_dashboard"))

        stored_filename = (
            f"{uuid.uuid4().hex}.{get_extension(original_filename)}"
        )
        assignment_dir.mkdir(parents=True, exist_ok=True)
        new_saved_path = assignment_dir / stored_filename
        old_stored_filename = assignment.assignment_stored_filename

        assignment_file.save(new_saved_path)
        assignment.assignment_original_filename = original_filename
        assignment.assignment_stored_filename = stored_filename

    try:
        db.session.commit()
        # Chỉ xóa tệp đề bài cũ SAU khi commit DB thành công
        if old_stored_filename:
            old_path = assignment_dir / old_stored_filename
            old_path.unlink(missing_ok=True)
        flash(f"Đã cập nhật bài tập «{assignment.title}».", "success")
    except Exception:
        db.session.rollback()
        if new_saved_path and new_saved_path.exists():
            new_saved_path.unlink(missing_ok=True)
        current_app.logger.exception("Không thể cập nhật bài tập.")
        flash("Có lỗi khi cập nhật bài tập.", "error")

    return redirect(url_for("main.teacher_dashboard"))


# ── Nhóm 3.3: Xóa bài tập (cascade xóa bài nộp + tệp) ──

@main.route(
    "/teacher/assignments/<int:assignment_id>/delete", methods=["POST"]
)
@login_required
def delete_assignment(assignment_id):
    assignment = db.session.get(Assignment, assignment_id)
    if assignment is None:
        flash("Không tìm thấy bài tập.", "error")
        return redirect(url_for("main.teacher_dashboard"))

    upload_dir = Path(current_app.config["UPLOAD_FOLDER"])
    assignment_dir = Path(current_app.config["ASSIGNMENT_FOLDER"])

    files_to_remove = [
        upload_dir / sub.stored_filename for sub in assignment.submissions
    ]
    if assignment.assignment_stored_filename:
        files_to_remove.append(
            assignment_dir / assignment.assignment_stored_filename
        )

    try:
        for sub in assignment.submissions:
            db.session.delete(sub)

        title = assignment.title
        db.session.delete(assignment)
        db.session.commit()

        # Xóa tệp vật lý sau khi giao dịch CSDL đã hoàn tất
        for fpath in files_to_remove:
            fpath.unlink(missing_ok=True)

        flash(
            f"Đã xóa bài tập «{title}» và tất cả bài nộp liên quan.",
            "success",
        )
    except Exception:
        db.session.rollback()
        current_app.logger.exception("Không thể xóa bài tập.")
        flash("Có lỗi khi xóa bài tập.", "error")

    return redirect(url_for("main.teacher_dashboard"))


# ── Nhóm 3.3: Xóa bài nộp ──

@main.route(
    "/teacher/submissions/<int:submission_id>/delete", methods=["POST"]
)
@login_required
def delete_submission(submission_id):
    submission = db.session.get(Submission, submission_id)
    if submission is None:
        flash("Không tìm thấy bài nộp.", "error")
        return redirect(url_for("main.teacher_dashboard"))

    upload_dir = Path(current_app.config["UPLOAD_FOLDER"])
    file_path = upload_dir / submission.stored_filename

    try:
        db.session.delete(submission)
        db.session.commit()
        file_path.unlink(missing_ok=True)
        flash("Đã xóa bài nộp.", "success")
    except Exception:
        db.session.rollback()
        current_app.logger.exception("Không thể xóa bài nộp.")
        flash("Có lỗi khi xóa bài nộp.", "error")

    return redirect(url_for("main.teacher_dashboard"))


# =====================================================================
#  TẢI TỆP
# =====================================================================

@main.route("/teacher/submissions/<int:submission_id>/download")
@login_required
def download_submission(submission_id):
    submission = db.session.get(Submission, submission_id)

    if submission is None:
        flash("Không tìm thấy bài nộp.", "error")
        return redirect(url_for("main.teacher_dashboard"))

    file_path = (
        Path(current_app.config["UPLOAD_FOLDER"]) / submission.stored_filename
    )
    if not file_path.is_file():
        flash("Tệp bài nộp không còn tồn tại trên máy chủ.", "error")
        return redirect(url_for("main.teacher_dashboard"))

    return send_from_directory(
        current_app.config["UPLOAD_FOLDER"],
        submission.stored_filename,
        as_attachment=True,
        download_name=submission.original_filename,
    )


@main.route("/assignments/<int:assignment_id>/download")
def download_assignment(assignment_id):
    assignment = db.session.get(Assignment, assignment_id)

    if assignment is None or not assignment.assignment_stored_filename:
        flash("Không tìm thấy đề bài.", "error")
        return redirect(url_for("main.index"))

    # Nhóm 1.7: giảng viên đã đăng nhập được tải đề bài kể cả khi đã đóng / quá hạn
    if not current_user.is_authenticated:
        if not assignment.is_open:
            flash("Bài tập đã đóng.", "error")
            return redirect(url_for("main.index"))
        if assignment.deadline and datetime.now() > assignment.deadline:
            flash("Bài tập đã quá hạn.", "error")
            return redirect(url_for("main.index"))

    file_path = (
        Path(current_app.config["ASSIGNMENT_FOLDER"])
        / assignment.assignment_stored_filename
    )
    if not file_path.is_file():
        flash("Tệp đề bài không còn tồn tại trên máy chủ.", "error")
        target = (
            "main.teacher_dashboard"
            if current_user.is_authenticated
            else "main.index"
        )
        return redirect(url_for(target))

    return send_from_directory(
        current_app.config["ASSIGNMENT_FOLDER"],
        assignment.assignment_stored_filename,
        as_attachment=True,
        download_name=assignment.assignment_original_filename,
    )


# =====================================================================
#  NHÓM 3.1: XUẤT EXCEL
# =====================================================================

@main.route("/teacher/export-excel")
@login_required
def export_excel():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    filter_assignment_id = request.args.get("assignment_id", 0, type=int)

    query = Submission.query.options(joinedload(Submission.assignment))
    if filter_assignment_id:
        query = query.filter(Submission.assignment_id == filter_assignment_id)

    submissions = query.order_by(Submission.submitted_at.desc()).all()

    latest_ids = _latest_submission_ids(filter_assignment_id or None)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bài nộp"

    headers = [
        "Mã nhận", "Bài tập", "Họ tên", "Số báo danh",
        "Mã sinh viên", "Tệp bài làm", "Thời gian nộp", "Trạng thái",
    ]
    ws.append(headers)

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(
        start_color="173B70", end_color="173B70", fill_type="solid"
    )
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for sub in submissions:
        ws.append([
            sub.id,
            _sanitize_excel_value(sub.assignment.title),
            _sanitize_excel_value(sub.student_name),
            _sanitize_excel_value(sub.student_number),
            _sanitize_excel_value(sub.student_code),
            _sanitize_excel_value(sub.original_filename),
            sub.submitted_at.strftime("%d/%m/%Y %H:%M:%S"),
            "Mới nhất" if sub.id in latest_ids else "Bản cũ",
        ])

    # Tự động điều chỉnh chiều rộng cột
    for col in ws.columns:
        max_length = 0
        for cell in col:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        adjusted = min(max_length + 3, 50)
        ws.column_dimensions[col[0].column_letter].width = adjusted

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    download_name = "bai_nop"
    if filter_assignment_id:
        a = db.session.get(Assignment, filter_assignment_id)
        if a:
            download_name = f"bai_nop_{a.title}"

    return send_file(
        buf,
        mimetype=(
            "application/vnd.openxmlformats-officedocument"
            ".spreadsheetml.sheet"
        ),
        as_attachment=True,
        download_name=f"{download_name}.xlsx",
    )


# =====================================================================
#  NHÓM 3.2: TẢI TẤT CẢ BÀI NỘP THÀNH ZIP
# =====================================================================

@main.route("/teacher/assignments/<int:assignment_id>/download-all")
@login_required
def download_all_submissions(assignment_id):
    assignment = db.session.get(Assignment, assignment_id)
    if assignment is None:
        flash("Không tìm thấy bài tập.", "error")
        return redirect(url_for("main.teacher_dashboard"))

    latest_only = request.args.get("latest_only", "0") == "1"

    query = Submission.query.filter_by(assignment_id=assignment_id)

    if latest_only:
        latest_select = (
            db.select(func.max(Submission.id))
            .filter(Submission.assignment_id == assignment_id)
            .group_by(Submission.student_code)
        )
        query = query.filter(Submission.id.in_(latest_select))

    submissions = query.order_by(Submission.student_code).all()

    if not submissions:
        flash("Chưa có bài nộp nào cho bài tập này.", "error")
        return redirect(url_for("main.teacher_dashboard"))

    upload_dir = Path(current_app.config["UPLOAD_FOLDER"])
    # Dùng SpooledTemporaryFile (tràn xuống đĩa nếu gói ZIP > 10 MB) để bảo vệ RAM
    buf = tempfile.SpooledTemporaryFile(max_size=10 * 1024 * 1024)

    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for sub in submissions:
            file_path = upload_dir / sub.stored_filename
            if not file_path.exists():
                continue

            if latest_only:
                # Mỗi sinh viên chỉ có 1 bài → dùng tên gốc
                arcname = f"{sub.student_code}/{sub.original_filename}"
            else:
                # Nhiều bài cùng sinh viên → thêm ID tránh trùng
                stem, dot, ext = sub.original_filename.rpartition(".")
                if dot:
                    arcname = (
                        f"{sub.student_code}/{stem}_{sub.id}.{ext}"
                    )
                else:
                    arcname = (
                        f"{sub.student_code}/"
                        f"{sub.original_filename}_{sub.id}"
                    )
            zf.write(file_path, arcname)

    buf.seek(0)

    suffix = "_moi_nhat" if latest_only else "_tat_ca"
    download_name = f"{assignment.title}{suffix}.zip"

    return send_file(
        buf,
        mimetype="application/zip",
        as_attachment=True,
        download_name=download_name,
    )